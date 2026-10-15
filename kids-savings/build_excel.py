"""Build Kids Savings.xlsx with live formulas (Excel tables auto-fill Balance on new rows)."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

from build_tracker import AVELYN, SAMUEL  # same verified data as the Numbers file

MONEY = '[Color10]$#,##0.00;[Red]-$#,##0.00;$0.00'
NAVY = "1E325A"
F = "Arial"
thin = Side(style="thin", color="BFBFBF")


def ledger(ws, name, headers, rows):
    ws.title = name
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    bal_letter = "ABCD"[headers.index("Balance")]
    for r, row in enumerate(rows, 2):
        ws.cell(r, 1, datetime.strptime(row[0], "%m/%d/%Y")).number_format = "m/d/yyyy"
        ws.cell(r, 2, float(row[1])).number_format = MONEY
        if "Note" in headers:
            ws.cell(r, 3, row[2] or None)
        ws[f"{bal_letter}{r}"] = f"=SUM(B$2:B{r})"
        ws[f"{bal_letter}{r}"].number_format = MONEY
    last = len(rows) + 1
    cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(headers)]
    cols[-1].calculatedColumnFormula = TableFormula(attr_text="SUM(B$2:B2)")
    t = Table(displayName=f"{name}Transactions", ref=f"A1:{bal_letter}{last}", tableColumns=cols)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for row in ws.iter_rows(min_row=1, max_row=last, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name=F, size=12, bold=cell.row == 1,
                             color="FFFFFF" if cell.row == 1 else "000000")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate([12, 12, 14, 12][: len(headers)]):
        ws.column_dimensions["ABCD"[i]].width = w

    # Headline total, outside the table so it never becomes part of it
    tc = "ABCDEF"[len(headers) + 1]
    ws[f"{tc}1"] = "Current Savings"
    ws[f"{tc}1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    ws[f"{tc}2"] = "=SUM(B:B)"
    ws[f"{tc}2"].font = Font(name=F, size=24, bold=True, color="1E823C")
    ws[f"{tc}2"].number_format = MONEY
    ws.column_dimensions[tc].width = 22
    ws.row_dimensions[2].height = 32
    ws[f"{tc}4"] = "To add: type the date and amount in the first empty row (use - for spending)."
    ws[f"{tc}4"].font = Font(name=F, size=10, italic=True, color="808080")
    ws.freeze_panes = "A2"
    return f"'{name}'!{tc}2"


wb = Workbook()
ov = wb.active
ov.title = "Overview"
a_ref = ledger(wb.create_sheet(), "Avelyn", ["Date", "Amount", "Balance"], AVELYN)
s_ref = ledger(wb.create_sheet(), "Samuel", ["Date", "Amount", "Note", "Balance"], SAMUEL)

ov["A1"], ov["B1"] = "Savings", "Balance"
ov["A2"], ov["B2"] = "Avelyn", f"={a_ref}"
ov["A3"], ov["B3"] = "Samuel", f"={s_ref}"
ov["A4"], ov["B4"] = "Combined Total", "=B2+B3"
for r in range(1, 5):
    for c in "AB":
        cell = ov[f"{c}{r}"]
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if r == 1:
            cell.font = Font(name=F, size=13, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
        else:
            cell.font = Font(name=F, size=18 if c == "B" else 16, bold=True,
                             color="1E823C" if c == "B" else NAVY)
    if r > 1:
        ov[f"B{r}"].number_format = MONEY
        ov.row_dimensions[r].height = 30
ov.column_dimensions["A"].width = 22
ov.column_dimensions["B"].width = 18

wb.calculation.fullCalcOnLoad = True
wb.save(Path(__file__).with_name("Kids Savings.xlsx"))
